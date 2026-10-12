"""Unit tests for transcript parsers and text cleaning (BE-5 & Core)."""

import io

import docx
import pytest

from app.parsers import Segment, TranscriptParseError, parse_transcript, to_canonical_text
from app.parsers.clean import clean_transcript_data


def test_clean_transcript_data_removes_bom_and_normalizes_crlf():
    raw_with_bom = b"\xef\xbb\xbfLine 1\r\nLine 2\rLine 3\n"
    cleaned = clean_transcript_data(raw_with_bom)
    assert cleaned == "Line 1\nLine 2\nLine 3\n"


def test_clean_transcript_data_invalid_encoding_raises():
    bad_bytes = b"\xff\xfe\x00\x00"  # Invalid UTF-8
    with pytest.raises(TranscriptParseError, match="File must be UTF-8 text"):
        clean_transcript_data(bad_bytes)


def test_parse_pasted_text_with_speakers_and_timestamps():
    raw = (
        "[00:01:05] Priya: Let's start the demo.\n"
        "[00:01:10] Rahul: I will fix the login bug by Wednesday.\n"
        "[00:01:25] Aman: I'll prepare onboarding.\n"
    )
    segments = parse_transcript(None, raw)
    assert len(segments) == 3
    assert segments[0].speaker_name == "Priya"
    assert segments[0].start_ms == 65000
    assert segments[0].end_ms == 70000
    assert segments[0].text == "Let's start the demo."

    assert segments[1].speaker_name == "Rahul"
    assert segments[1].start_ms == 70000
    assert segments[1].end_ms == 85000

    assert segments[2].speaker_name == "Aman"
    assert segments[2].start_ms == 85000
    assert segments[2].end_ms == 85000


def test_parse_text_colon_in_time_not_treated_as_speaker():
    # "Meeting at 10:30" should NOT treat "Meeting at 10" as speaker
    raw = "Meeting at 10:30 is scheduled for everyone."
    segments = parse_transcript(None, raw)
    assert len(segments) == 1
    assert segments[0].speaker_name is None
    assert segments[0].text == "Meeting at 10:30 is scheduled for everyone."


def test_parse_vtt_format():
    vtt_content = (
        "WEBVTT\n\n"
        "NOTE This is an internal meeting note\n\n"
        "00:01:02.345 --> 00:01:05.000\n"
        "<v Priya>We ship Friday.\n\n"
        "00:01:06.000 --> 00:01:09.500\n"
        "<v Rahul>Sounds good.\n"
    )
    segments = parse_transcript("meeting.vtt", vtt_content)
    assert len(segments) == 2
    assert segments[0].speaker_name == "Priya"
    assert segments[0].text == "We ship Friday."
    assert segments[0].start_ms == 62345
    assert segments[0].end_ms == 65000

    assert segments[1].speaker_name == "Rahul"
    assert segments[1].text == "Sounds good."


def test_parse_srt_format():
    srt_content = (
        "1\n"
        "00:01:02,345 --> 00:01:05,000\n"
        "Priya: We ship Friday.\n\n"
        "2\n"
        "00:01:06,000 --> 00:01:09,000\n"
        "Rahul: Roger that.\n"
    )
    segments = parse_transcript("meeting.srt", srt_content)
    assert len(segments) == 2
    assert segments[0].speaker_name == "Priya"
    assert segments[0].text == "We ship Friday."
    assert segments[0].start_ms == 62345
    assert segments[0].end_ms == 65000

    assert segments[1].speaker_name == "Rahul"
    assert segments[1].text == "Roger that."


def test_parse_docx_format():
    doc = docx.Document()
    doc.add_paragraph("Priya: Welcome everyone.")
    doc.add_paragraph("Rahul: I'll complete the task.")

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    segments = parse_transcript("meeting.docx", docx_bytes)
    assert len(segments) == 2
    assert segments[0].speaker_name == "Priya"
    assert segments[0].text == "Welcome everyone."
    assert segments[1].speaker_name == "Rahul"
    assert segments[1].text == "I'll complete the task."


def test_parse_transcript_unsupported_extension_raises():
    with pytest.raises(TranscriptParseError, match="Unsupported transcript format"):
        parse_transcript("notes.mp3", b"audio data")


def test_parse_transcript_empty_segments_raises():
    with pytest.raises(TranscriptParseError, match="no valid text segments"):
        parse_transcript("empty.txt", "   \n\n  ")


def test_to_canonical_text():
    segments = [
        Segment(speaker_name="Priya", start_ms=0, end_ms=1000, text="First item."),
        Segment(speaker_name=None, start_ms=None, end_ms=None, text="Anonymous statement."),
        Segment(speaker_name="Rahul", start_ms=2000, end_ms=3000, text="Final confirmation."),
    ]
    canonical = to_canonical_text(segments)
    expected = (
        "Priya: First item.\n"
        "Anonymous statement.\n"
        "Rahul: Final confirmation."
    )
    assert canonical == expected
